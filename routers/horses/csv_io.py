import csv
import io
from openpyxl import load_workbook
from fastapi import APIRouter, Request, UploadFile, File
from fastapi.responses import RedirectResponse, StreamingResponse
from sqlmodel import Session, select

from database import engine
from models import Horse
from templates import templates

router = APIRouter()

@router.get("/horses/export/{race_id}")
def export_horses(race_id: int):
    with Session(engine) as session:
        horses = session.exec(select(Horse).where(Horse.race_id == race_id)).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "馬名", "脚質", "メモタグ", "過去走内容スコア", "コース適性スコア",
        "展開スコア", "オッズ",
    ])
    for h in horses:
        writer.writerow([
            h.name, h.running_style, h.memo_tag,
            h.past_performance_score, h.course_aptitude_score,
            h.pace_score, h.odds,
        ])

    csv_data = "\ufeff" + output.getvalue()
    return StreamingResponse(
        iter([csv_data]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=race_{race_id}_horses.csv"},
    )

@router.post("/horses/import/{race_id}")
async def import_horses(race_id: int, file: UploadFile = File(...)):
    content = await file.read()
    rows = []

    if file.filename and file.filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
    else:
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = content.decode("cp932")
        reader = csv.reader(io.StringIO(decoded))
        next(reader, None)
        rows = list(reader)

    with Session(engine) as session:
        for row in rows:
            if not row or not row[0]:
                continue
            row = (list(row) + [""] * 7)[:7]
            name, running_style, memo_tag, pp, ca, pace, odds = row
            horse = Horse(
                race_id=race_id,
                name=name,
                running_style=running_style,
                memo_tag=memo_tag,
                past_performance_score=float(pp) if pp else 0.0,
                course_aptitude_score=float(ca) if ca else 0.0,
                pace_score=float(pace) if pace else 0.0,
                odds=float(odds) if odds else None,
            )
            session.add(horse)
        session.commit()
    return RedirectResponse(url=f"/horses/ranking-page/{race_id}", status_code=303)

@router.get("/horses/import/{race_id}")
def import_horses_page(request: Request, race_id: int):
    return templates.TemplateResponse("horse_import.html", {"request": request, "race_id": race_id})

@router.get("/horses/import-template")
def download_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "馬名", "脚質", "メモタグ", "過去走内容スコア", "コース適性スコア",
        "展開スコア", "オッズ",
    ])

    csv_data = "\ufeff" + output.getvalue()
    return StreamingResponse(
        iter([csv_data]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=import_template.csv"},
    )