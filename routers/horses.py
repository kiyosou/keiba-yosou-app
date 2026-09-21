import csv
import io
from openpyxl import load_workbook
from fastapi import APIRouter, Request, Form, UploadFile, File, Query
from fastapi.responses import RedirectResponse, StreamingResponse
from sqlmodel import Session, select

from database import engine
from models import Race, Horse
from scoring import race_label, calculate_score
from templates import templates
from parsing import parse_shutsuba_text

router = APIRouter()

def get_all_horse_names(session: Session) -> list[str]:
    horses = session.exec(select(Horse)).all()
    return sorted(set(h.name for h in horses))

@router.get("/horses/new")
def horse_add_page(request: Request):
    with Session(engine) as session:
        races = session.exec(select(Race)).all()
        horse_names = get_all_horse_names(session)
    race_options = [{"id": r.id, "label": race_label(r)} for r in races]
    return templates.TemplateResponse("input.html", {
        "request": request, "race_options": race_options, "horse_names": horse_names,
    })
@router.post("/horses/add")
def add_horse(
    race_id: int = Form(...),
    name: str = Form(...),
    running_style: str = Form(""),
    memo_tag: str = Form(""),
    past_performance_score: float = Form(...),
    course_aptitude_score: float = Form(...),
    pace_score: float = Form(...),
    odds: float = Form(None),
):
    horse = Horse(
        race_id=race_id, name=name,
        running_style=running_style, memo_tag=memo_tag,
        past_performance_score=past_performance_score,
        course_aptitude_score=course_aptitude_score,
        pace_score=pace_score,
        odds=odds,
    )
    with Session(engine) as session:
        session.add(horse)
        session.commit()
    return RedirectResponse(url="/", status_code=303)

@router.get("/horses/ranking-page/{race_id}")
def ranking_page(
    request: Request, race_id: int,
    back_date: str = "", venue: str = "", surface: str = "",
    min_distance: str = "", max_distance: str = "",
):
    with Session(engine) as session:
        race = session.get(Race, race_id)
        horses = session.exec(select(Horse).where(Horse.race_id == race_id)).all()
        ranked = sorted(horses, key=lambda h: calculate_score(h), reverse=True)
        horse_list = [
            {
                "id": h.id,
                "name": h.name,
                "score": calculate_score(h),
                "odds": h.odds,
                "running_style": h.running_style,
                "memo_tag": h.memo_tag,
            }
            for h in ranked
        ]

    back_url = ""
    if back_date:
        params = []
        if venue:
            params.append(f"venue={venue}")
        if surface:
            params.append(f"surface={surface}")
        if min_distance:
            params.append(f"min_distance={min_distance}")
        if max_distance:
            params.append(f"max_distance={max_distance}")
        back_url = f"/races/{back_date}" + ("?" + "&".join(params) if params else "")

    return templates.TemplateResponse("ranking.html", {
        "request": request, "horses": horse_list,
        "race_label": race_label(race) if race else "",
        "race_id": race_id, "back_url": back_url,
    })
@router.get("/horses/export/{race_id}")
def export_horses(race_id: int):
    with Session(engine) as session:
        horses = session.exec(select(Horse).where(Horse.race_id == race_id)).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "馬名", "脚質", "メモタグ", "過去走内容スコア", "コース適性スコア",
        "展開スコア", "オッズ",
    ])
    for h in horses:
        writer.writerow([
            h.name, h.running_style, h.memo_tag,
            h.past_performance_score, h.course_aptitude_score,
            h.pace_score, h.odds,
        ])

    csv_data = "\ufeff" + output.getvalue()
    return StreamingResponse(
        iter([csv_data]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=race_{race_id}_horses.csv"},
    )

@router.post("/horses/import/{race_id}")
async def import_horses(race_id: int, file: UploadFile = File(...)):
    content = await file.read()
    rows = []

    if file.filename and file.filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
    else:
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = content.decode("cp932")
        reader = csv.reader(io.StringIO(decoded))
        next(reader, None)
        rows = list(reader)

    with Session(engine) as session:
        for row in rows:
            if not row or not row[0]:
                continue
            row = (list(row) + [""] * 7)[:7]
            name, running_style, memo_tag, pp, ca, pace, odds = row
            horse = Horse(
                race_id=race_id,
                name=name,
                running_style=running_style,
                memo_tag=memo_tag,
                past_performance_score=float(pp) if pp else 0.0,
                course_aptitude_score=float(ca) if ca else 0.0,
                pace_score=float(pace) if pace else 0.0,
                odds=float(odds) if odds else None,
            )
            session.add(horse)
        session.commit()
    return RedirectResponse(url=f"/horses/ranking-page/{race_id}", status_code=303)

@router.get("/horses/import/{race_id}")
def import_horses_page(request: Request, race_id: int):
    return templates.TemplateResponse("horse_import.html", {"request": request, "race_id": race_id})



@router.get("/horses/import-template")
def download_import_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "馬名", "脚質", "メモタグ", "過去走内容スコア", "コース適性スコア",
        "展開スコア", "オッズ",
    ])

    csv_data = "\ufeff" + output.getvalue()
    return StreamingResponse(
        iter([csv_data]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=import_template.csv"},
    )

@router.post("/horses/{horse_id}/delete")
def delete_horse(horse_id: int):
    with Session(engine) as session:
        horse = session.get(Horse, horse_id)
        if horse:
            race_id = horse.race_id
            session.delete(horse)
            session.commit()
        else:
            race_id = None
    return RedirectResponse(url=f"/horses/ranking-page/{race_id}", status_code=303)

@router.get("/horses/profile/{horse_name}")
def horse_profile(request: Request, horse_name: str):
    with Session(engine) as session:
        horses = session.exec(select(Horse).where(Horse.name == horse_name)).all()
        history = []
        for h in horses:
            race = session.get(Race, h.race_id)
            history.append({
                "race_id": h.race_id,
                "race_label": race_label(race) if race else "不明なレース",
                "race_date": race.date if race else "",
                "running_style": h.running_style,
                "memo_tag": h.memo_tag,
                "score": calculate_score(h),
                "odds": h.odds,
            })
        history.sort(key=lambda x: x["race_date"], reverse=True)

    return templates.TemplateResponse("horse_profile.html", {
        "request": request, "horse_name": horse_name, "history": history,
    })
@router.get("/horses/paste-import/{race_id}")
def paste_import_page(request: Request, race_id: int):
    return templates.TemplateResponse("horse_paste_import.html", {
        "request": request, "race_id": race_id, "entries": None, "raw_text": "",
    })

@router.post("/horses/paste-import/{race_id}")
def paste_import_parse(request: Request, race_id: int, raw_text: str = Form(...)):
    entries = parse_shutsuba_text(raw_text)
    return templates.TemplateResponse("horse_paste_import.html", {
        "request": request, "race_id": race_id, "entries": entries, "raw_text": raw_text,
    })

@router.post("/horses/bulk-add/{race_id}")
async def bulk_add_horses(race_id: int, request: Request):
    form = await request.form()
    count = int(form.get("count", 0))

    with Session(engine) as session:
        for i in range(count):
            name = form.get(f"name_{i}")
            if not name:
                continue
            odds_raw = form.get(f"odds_{i}") or ""
            horse = Horse(
                race_id=race_id,
                name=name,
                running_style=form.get(f"running_style_{i}", ""),
                memo_tag=form.get(f"memo_tag_{i}", ""),
                past_performance_score=float(form.get(f"past_performance_score_{i}") or 0),
                course_aptitude_score=float(form.get(f"course_aptitude_score_{i}") or 0),
                pace_score=float(form.get(f"pace_score_{i}") or 0),
                odds=float(odds_raw) if odds_raw else None,
            )
            session.add(horse)
        session.commit()
    return RedirectResponse(url=f"/horses/ranking-page/{race_id}", status_code=303)

@router.get("/horses/{horse_id}/edit")
def horse_edit_page(request: Request, horse_id: int):
    with Session(engine) as session:
        horse = session.get(Horse, horse_id)
        races = session.exec(select(Race)).all()
    race_options = [{"id": r.id, "label": race_label(r)} for r in races]
    return templates.TemplateResponse("horse_edit.html", {
        "request": request, "horse": horse, "race_options": race_options,
    })

@router.post("/horses/{horse_id}/edit")
def horse_edit_submit(
    horse_id: int,
    race_id: int = Form(...),
    name: str = Form(...),
    running_style: str = Form(""),
    memo_tag: str = Form(""),
    past_performance_score: float = Form(...),
    course_aptitude_score: float = Form(...),
    pace_score: float = Form(...),
    odds: float = Form(None),
):
    with Session(engine) as session:
        horse = session.get(Horse, horse_id)
        if horse:
            horse.race_id = race_id
            horse.name = name
            horse.running_style = running_style
            horse.memo_tag = memo_tag
            horse.past_performance_score = past_performance_score
            horse.course_aptitude_score = course_aptitude_score
            horse.pace_score = pace_score
            horse.odds = odds
            session.add(horse)
            session.commit()
    return RedirectResponse(url=f"/horses/ranking-page/{race_id}", status_code=303)

@router.get("/horses")
def horses_list(request: Request, style: str = ""):
    with Session(engine) as session:
        horses = session.exec(select(Horse)).all()

        by_name = {}
        for h in horses:
            race = session.get(Race, h.race_id)
            race_date = race.date if race else ""
            entry = {
                "id": h.id,
                "name": h.name,
                "race_id": h.race_id,
                "race_label": race_label(race) if race else "不明なレース",
                "race_date": race_date,
                "running_style": h.running_style,
                "memo_tag": h.memo_tag,
                "score": calculate_score(h),
                "odds": h.odds,
            }
            # 同じ馬名の中で、race_dateが一番新しいものだけを残す
            if h.name not in by_name or race_date > by_name[h.name]["race_date"]:
                by_name[h.name] = entry

        horse_list = list(by_name.values())
        if style:
            horse_list = [h for h in horse_list if h["running_style"] == style]

    horse_list.sort(key=lambda x: x["race_date"], reverse=True)
    return templates.TemplateResponse("horses_list.html", {
        "request": request, "horses": horse_list, "style": style,
    })

@router.get("/horses/compare")
def horses_compare(request: Request, names: list[str] = Query(default=[])):
    if len(names) > 18:
        names = names[:18]

    with Session(engine) as session:
        horse_histories = []
        for name in names:
            horses = session.exec(select(Horse).where(Horse.name == name)).all()
            history = []
            for h in horses:
                race = session.get(Race, h.race_id)
                history.append({
                    "race_id": h.race_id,
                    "race_label": race_label(race) if race else "不明なレース",
                    "race_date": race.date if race else "",
                    "running_style": h.running_style,
                    "memo_tag": h.memo_tag,
                    "score": calculate_score(h),
                    "odds": h.odds,
                })
            history.sort(key=lambda x: x["race_date"], reverse=True)
            horse_histories.append({"name": name, "history": history})

    return templates.TemplateResponse("horses_compare.html", {
        "request": request, "horse_histories": horse_histories,
    })