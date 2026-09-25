import io
import re
import unicodedata
import threading
from fastapi import APIRouter, Request, Form, UploadFile, File
from fastapi.responses import RedirectResponse
from openpyxl import load_workbook

from sqlmodel import Session, select
from database import engine
from models import StandardTime
from templates import templates
from parsing import time_str_to_seconds

router = APIRouter()

VENUES = ["札幌", "函館", "福島", "新潟", "東京", "中山", "中京", "京都", "阪神", "小倉"]

# 進捗状況を保存する場所(シンプルにメモリ上で管理)
import_progress = {"total": 0, "done": 0, "finished": False, "imported_count": 0}

def run_import(venue: str, content: bytes):
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)

        total_rows = 0
        valid_sheets = []
        for sheet_name in workbook.sheetnames:
            if sheet_name.startswith("芝") or sheet_name.startswith("ダ"):
                rest = sheet_name[1:]
                if not re.match(r"^(\d+)(外|内)?$", rest):
                    continue
                sheet = workbook[sheet_name]
                valid_sheets.append(sheet_name)
                total_rows += sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if row and row[0])

        import_progress["total"] = total_rows
        import_progress["done"] = 0
        import_progress["finished"] = False
        import_progress["imported_count"] = 0
        import_progress["error"] = ""
        print(f"[基準タイム取り込み] 開始: 全{total_rows}件")

        new_records = []

        with Session(engine) as session:
            # 既存データを先にまとめて削除(1件ずつの確認をなくす)
            existing_all = session.exec(
                select(StandardTime).where(StandardTime.venue == venue)
            ).all()
            for ex in existing_all:
                session.delete(ex)
            session.commit()
            print(f"[基準タイム取り込み] 既存{len(existing_all)}件を削除しました")

            for sheet_name in valid_sheets:
                sheet = workbook[sheet_name]
                surface = "芝" if sheet_name.startswith("芝") else "ダート"
                rest = sheet_name[1:]
                m = re.match(r"^(\d+)(外|内)?$", rest)
                if not m:
                    continue
                distance = int(m.group(1))
                course_type = m.group(2) or ""

                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    age, race_class = row[0], row[1]
                    age = unicodedata.normalize("NFKC", str(age)).strip() if age else age
                    race_class = unicodedata.normalize("NFKC", str(race_class)).strip() if race_class else race_class
                    if not age or not race_class:
                        import_progress["done"] += 1
                        continue

                    race_count = row[2] or 0
                    winner_time = str(row[4]) if row[4] else ""
                    top3_avg_time = str(row[7]) if row[7] else ""
                    pci3 = row[10] or 0
                    ave_3f = row[11] or 0

                    st = StandardTime(
                        venue=venue, surface=surface, distance=distance, course_type=course_type,
                        age=age, race_class=race_class,
                        race_count=int(race_count) if race_count else 0,
                        winner_time=winner_time,
                        winner_time_seconds=time_str_to_seconds(winner_time),
                        top3_avg_time=top3_avg_time,
                        top3_avg_seconds=time_str_to_seconds(top3_avg_time),
                        pci3=float(pci3) if pci3 else 0.0,
                        ave_3f=float(ave_3f) if ave_3f else 0.0,
                    )
                    new_records.append(st)
                    import_progress["done"] += 1
                    if import_progress["done"] % 20 == 0:
                        print(f"[基準タイム取り込み] {import_progress['done']}/{total_rows}件 読み込み済み")

            # ここでまとめて1回の通信で保存する
            print(f"[基準タイム取り込み] {len(new_records)}件をまとめて保存しています...")
            session.bulk_save_objects(new_records)
            session.commit()

        import_progress["imported_count"] = len(new_records)
        import_progress["finished"] = True
        print(f"[基準タイム取り込み] 完了: {len(new_records)}件")

    except Exception as e:
        import traceback
        traceback.print_exc()
        import_progress["error"] = str(e)
        import_progress["finished"] = True
@router.get("/standard-times/import")
def standard_time_import_page(request: Request):
    return templates.TemplateResponse("standard_time_import.html", {
        "request": request, "venues": VENUES,
    })

@router.post("/standard-times/import")
async def standard_time_import(venue: str = Form(...), file: UploadFile = File(...)):
    content = await file.read()
    thread = threading.Thread(target=run_import, args=(venue, content))
    thread.start()
    return RedirectResponse(url="/standard-times/import/progress", status_code=303)

@router.get("/standard-times/import/progress")
def import_progress_page(request: Request):
    return templates.TemplateResponse("standard_time_progress.html", {"request": request})

@router.get("/standard-times/import/status")
def import_status():
    return import_progress

AGE_ORDER = ["2歳", "3歳", "古馬", "全"]
CLASS_ORDER = ["新馬", "未勝利", "500万", "1000万", "1600万", "OPEN", "平均等"]

def sort_key(t):
    age = unicodedata.normalize("NFKC", t.age or "").strip()
    race_class = unicodedata.normalize("NFKC", t.race_class or "").strip()
    age_index = AGE_ORDER.index(age) if age in AGE_ORDER else len(AGE_ORDER)
    class_index = CLASS_ORDER.index(race_class) if race_class in CLASS_ORDER else len(CLASS_ORDER)
    return (t.venue, t.surface, t.distance, t.course_type, age_index, class_index)

@router.get("/standard-times")
def standard_times_list(request: Request, venue: str = "", distance: str = ""):
    with Session(engine) as session:
        times = session.exec(select(StandardTime)).all()

    all_venues = sorted(set(t.venue for t in times))
    all_distances = sorted(set(t.distance for t in times))

    filtered = times
    if venue:
        filtered = [t for t in filtered if t.venue == venue]
    if distance:
        filtered = [t for t in filtered if t.distance == int(distance)]

    times_sorted = sorted(filtered, key=sort_key)

    return templates.TemplateResponse("standard_times.html", {
        "request": request, "times": times_sorted,
        "venues": all_venues, "distances": all_distances,
        "selected_venue": venue, "selected_distance": distance,
    })